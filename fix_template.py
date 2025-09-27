import os
import re
from docx import Document
from pathlib import Path

def fix_placeholder_in_template(template_path):
    """Sửa placeholder có khoảng trắng trong template Word"""
    try:
        doc = Document(template_path)
        
        # Tìm và thay thế placeholder có khoảng trắng
        changes_made = 0
        
        # Xử lý paragraphs
        for para in doc.paragraphs:
            if '{{Ngay_ ĐG}}' in para.text:
                para.text = para.text.replace('{{Ngay_ ĐG}}', '{{Ngay_ĐG}}')
                changes_made += 1
                print(f"   - Đã sửa placeholder trong paragraph: {para.text[:50]}...")
        
        # Xử lý tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if '{{Ngay_ ĐG}}' in para.text:
                            para.text = para.text.replace('{{Ngay_ ĐG}}', '{{Ngay_ĐG}}')
                            changes_made += 1
                            print(f"   - Đã sửa placeholder trong table: {para.text[:50]}...")
        
        # Xử lý headers và footers
        for section in doc.sections:
            if section.header:
                for para in section.header.paragraphs:
                    if '{{Ngay_ ĐG}}' in para.text:
                        para.text = para.text.replace('{{Ngay_ ĐG}}', '{{Ngay_ĐG}}')
                        changes_made += 1
                        print(f"   - Đã sửa placeholder trong header: {para.text[:50]}...")
            
            if section.footer:
                for para in section.footer.paragraphs:
                    if '{{Ngay_ ĐG}}' in para.text:
                        para.text = para.text.replace('{{Ngay_ ĐG}}', '{{Ngay_ĐG}}')
                        changes_made += 1
                        print(f"   - Đã sửa placeholder trong footer: {para.text[:50]}...")
        
        if changes_made > 0:
            # Lưu file đã sửa
            doc.save(template_path)
            print(f"✅ Đã sửa {changes_made} placeholder trong {template_path.name}")
            return True
        else:
            print(f"ℹ️ Không tìm thấy placeholder cần sửa trong {template_path.name}")
            return False
            
    except Exception as e:
        print(f"❌ Lỗi khi sửa template {template_path.name}: {e}")
        return False

def main():
    print("🔧 Bắt đầu sửa template có placeholder lỗi...")
    print("=" * 60)
    
    # Sửa template "7. Ke hoach danh gia.docx"
    template_path = Path("templates/7. Ke hoach danh gia.docx")
    
    if template_path.exists():
        print(f"📝 Đang sửa template: {template_path.name}")
        success = fix_placeholder_in_template(template_path)
        
        if success:
            print("✅ Hoàn thành sửa template!")
        else:
            print("⚠️ Không cần sửa hoặc có lỗi xảy ra")
    else:
        print(f"❌ Không tìm thấy file: {template_path}")
    
    print("\n" + "=" * 60)
    print("🎯 Bây giờ có thể chạy lại script run.py")

if __name__ == "__main__":
    main()
