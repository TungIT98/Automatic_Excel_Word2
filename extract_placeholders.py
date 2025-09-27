import os
import re
from docx import Document
from pathlib import Path

def extract_placeholders_from_docx(file_path):
    """Trích xuất tất cả placeholder từ file Word"""
    try:
        doc = Document(file_path)
        placeholders = set()
        
        # Tìm placeholder trong paragraphs
        for para in doc.paragraphs:
            matches = re.findall(r"\{\{(.*?)\}\}", para.text)
            placeholders.update(matches)
        
        # Tìm placeholder trong tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        matches = re.findall(r"\{\{(.*?)\}\}", para.text)
                        placeholders.update(matches)
        
        # Tìm placeholder trong headers và footers
        for section in doc.sections:
            # Header
            if section.header:
                for para in section.header.paragraphs:
                    matches = re.findall(r"\{\{(.*?)\}\}", para.text)
                    placeholders.update(matches)
            
            # Footer
            if section.footer:
                for para in section.footer.paragraphs:
                    matches = re.findall(r"\{\{(.*?)\}\}", para.text)
                    placeholders.update(matches)
        
        return sorted(list(placeholders))
    except Exception as e:
        print(f"❌ Lỗi khi đọc file {file_path}: {e}")
        return []

def scan_folder_for_placeholders(folder_path):
    """Quét tất cả file Word trong thư mục và trích xuất placeholder"""
    results = {}
    folder = Path(folder_path)
    
    if not folder.exists():
        print(f"❌ Thư mục không tồn tại: {folder_path}")
        return results
    
    docx_files = list(folder.glob("*.docx"))
    if not docx_files:
        print(f"❌ Không tìm thấy file .docx nào trong {folder_path}")
        return results
    
    print(f"🔍 Đang quét {len(docx_files)} file Word...")
    
    for file_path in docx_files:
        print(f"📖 Đang xử lý: {file_path.name}")
        placeholders = extract_placeholders_from_docx(file_path)
        results[file_path.name] = placeholders
    
    return results

def generate_template_config(placeholder_map):
    """Tạo cấu hình template từ placeholder map"""
    config = {}
    
    for filename, placeholders in placeholder_map.items():
        if placeholders:
            config[filename] = {}
            for placeholder in placeholders:
                # Tạo giá trị mẫu cho placeholder
                if "ten" in placeholder.lower() or "name" in placeholder.lower():
                    config[filename][placeholder] = "Tên mẫu"
                elif "dia_chi" in placeholder.lower() or "address" in placeholder.lower():
                    config[filename][placeholder] = "Địa chỉ mẫu"
                elif "ngay" in placeholder.lower() or "date" in placeholder.lower() or "time" in placeholder.lower():
                    config[filename][placeholder] = "27/09/2025"
                elif "so" in placeholder.lower() or "number" in placeholder.lower():
                    config[filename][placeholder] = "001/2025"
                elif "email" in placeholder.lower():
                    config[filename][placeholder] = "example@email.com"
                elif "phone" in placeholder.lower() or "sdt" in placeholder.lower():
                    config[filename][placeholder] = "0123456789"
                else:
                    config[filename][placeholder] = f"Giá trị cho {placeholder}"
    
    return config

if __name__ == "__main__":
    # Đường dẫn thư mục templates
    folder = r"C:\AutomaticExcel\templates"
    
    print("🚀 Bắt đầu quét placeholder từ các file Word...")
    print("=" * 60)
    
    # Quét placeholder
    placeholder_map = scan_folder_for_placeholders(folder)
    
    if not placeholder_map:
        print("❌ Không tìm thấy placeholder nào!")
        exit()
    
    print("\n📋 KẾT QUẢ QUÉT PLACEHOLDER:")
    print("=" * 60)
    
    total_placeholders = 0
    for filename, placeholders in placeholder_map.items():
        if placeholders:
            print(f"\n📄 {filename}:")
            for p in placeholders:
                print(f"   - {{ {p} }}")
            total_placeholders += len(placeholders)
        else:
            print(f"\n📄 {filename}: (Không có placeholder)")
    
    print(f"\n📊 TỔNG KẾT:")
    print(f"   - Số file đã quét: {len(placeholder_map)}")
    print(f"   - Tổng số placeholder: {total_placeholders}")
    
    # Tạo cấu hình template
    print(f"\n🔧 TẠO CẤU HÌNH TEMPLATE:")
    print("=" * 60)
    
    template_config = generate_template_config(placeholder_map)
    
    print("TEMPLATE_CONFIG = {")
    for filename, config in template_config.items():
        if config:
            print(f'    "{filename}": {{')
            for key, value in config.items():
                print(f'        "{key}": "{value}",')
            print("    },")
    print("}")
    
    print(f"\n✅ Hoàn thành! Đã tạo cấu hình cho {len(template_config)} template.")